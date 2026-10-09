"""分析結果・予想を Word（.docx）に書き出す。表の色分けは Excel 版（xlsx.py）に合わせる。"""
import unicodedata

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Yu Gothic"
HEADER_FILL = "1F4E78"
GOOD, BAD = "E2EFDA", "FCE4D6"
MARK_FILL = {"◎": "FFF2CC", "○": "E2EFDA"}
MUTED = RGBColor(0x59, 0x59, 0x59)


def _set_font(style_or_run, size=None, bold=None, color=None):
    f = style_or_run.font
    f.name = FONT
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = color
    rpr = style_or_run.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def new_document(landscape=True):
    doc = Document()
    # python-docx の既定テンプレートは w:zoom に必須の w:percent が無く、厳密な検査で弾かれる
    for z in doc.settings.element.findall(qn("w:zoom")):
        z.set(qn("w:percent"), "100")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))
    _set_font(doc.styles["Normal"], size=10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(3)
    for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
        _set_font(doc.styles[name], size=size, bold=True, color=RGBColor(0x1F, 0x4E, 0x78))
    for name in ("List Bullet",):
        _set_font(doc.styles[name], size=10)
    return doc


def text_width(doc):
    sec = doc.sections[0]
    return sec.page_width - sec.left_margin - sec.right_margin


def title(doc, text):
    doc.add_paragraph(text, style="Title")


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def para(doc, text, muted=False, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_font(r, size=9 if muted else 10, bold=bold, color=MUTED if muted else None)
    return p


def bullet(doc, text, label=None):
    p = doc.add_paragraph(style="List Bullet")
    if label:
        _set_font(p.add_run(f"{label}: "), bold=True)
    _set_font(p.add_run(text))
    return p


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _disp(text):
    """表示幅（全角=2、半角=1）。"""
    return sum(2 if unicodedata.east_asian_width(ch) in "WFA" else 1 for ch in text)


def _fmt(col, v):
    """セルの文字列と塗り色。"""
    if v is None or v == "":
        return ("" if col == "印" else "–"), None
    if isinstance(v, float):
        if "率" in col:
            return f"{v:.1%}", None
        if "A/E" in col:
            return f"{v:.2f}", GOOD if v >= 1.15 else BAD if v <= 0.85 else None
        return f"{v:.1f}", None
    if col == "印":
        return str(v), MARK_FILL.get(v)
    return str(v), None


def table(doc, rows, font_size=8.5):
    """rows（dict のリスト）を表にする。列幅は中身の長さから決めて、本文幅に収める。"""
    if not rows:
        para(doc, "データなし", muted=True)
        return
    cols = list(rows[0])
    cells = [[_fmt(c, r.get(c)) for c in cols] for r in rows]
    # 見出しは2行まで折り返してよいので半分の幅で数える
    weight = [max([_disp(c) / 2, 4] + [_disp(row[i][0]) for row in cells]) + 1.5 for i, c in enumerate(cols)]
    weight = [min(w, 44) for w in weight]
    total = text_width(doc)
    widths = [int(total * w / sum(weight)) for w in weight]

    t = doc.add_table(rows=1 + len(rows), cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, c in enumerate(cols):
        t.columns[i].width = widths[i]
    for j, values in enumerate([[(c, HEADER_FILL) for c in cols]] + cells):
        for i, (text, fill) in enumerate(values):
            cell = t.rows[j].cells[i]
            cell.width = widths[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            if j == 0:
                _set_font(run, size=font_size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                _set_font(run, size=font_size)
                raw = rows[j - 1].get(cols[i])
                if isinstance(raw, (int, float)) and not isinstance(raw, bool):
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if fill:
                _shade(cell, fill)
    # 見出し行をページごとに繰り返す
    trpr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trpr.append(hdr)
    doc.add_paragraph()
