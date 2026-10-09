"""分析結果・予想を Excel（.xlsx）に書き出す。書式は scripts/build_xlsx.py に合わせる。"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

F = "Arial"
hdr_font = Font(name=F, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E78")
base = Font(name=F, size=10)
title_font = Font(name=F, bold=True, size=12)
note_font = Font(name=F, size=9, italic=True, color="595959")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
good = PatternFill("solid", fgColor="E2EFDA")
bad = PatternFill("solid", fgColor="FCE4D6")
mark_fill = {"◎": PatternFill("solid", fgColor="FFF2CC"), "○": PatternFill("solid", fgColor="E2EFDA")}


def _fmt(cell, col, v):
    cell.font, cell.border = base, border
    if isinstance(v, float):
        if "率" in col:
            cell.number_format = "0.0%"
        elif "A/E" in col:
            cell.number_format = "0.00"
            cell.fill = good if v >= 1.15 else bad if v <= 0.85 else PatternFill()
        else:
            cell.number_format = "0.0"


def _width(col, values):
    n = max([len(str(col))] + [len(str(v)) for v in values if v is not None] or [4])
    return min(max(6, n * 1.8 + 2), 40)


def table_at(ws, row, rows, name=None, title=None, note=None):
    """ws の row 行目から表を書く。次に書ける行番号を返す。"""
    if title:
        ws.cell(row=row, column=1, value=title).font = title_font
        row += 1
    if note:
        ws.cell(row=row, column=1, value=note).font = note_font
        row += 1
    if not rows:
        ws.cell(row=row, column=1, value="データなし").font = note_font
        return row + 2
    cols = list(rows[0])
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font, x.fill, x.border = hdr_font, hdr_fill, border
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0,
                                                 _width(c, [r.get(c) for r in rows[:200]]))
    for j, r in enumerate(rows, row + 1):
        for i, c in enumerate(cols, 1):
            v = r.get(c)
            x = ws.cell(row=j, column=i, value=v)
            _fmt(x, c, v)
            if c == "印" and v in mark_fill:
                x.fill = mark_fill[v]
    if name:
        ref = f"A{row}:{get_column_letter(len(cols))}{row + len(rows)}"
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(t)
    return row + len(rows) + 2


def text_sheet(ws, lines):
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    for i, line in enumerate(lines, 1):
        if isinstance(line, tuple):
            ws.cell(row=i, column=1, value=line[0]).font = Font(name=F, bold=True, size=10)
            ws.cell(row=i, column=2, value=line[1]).font = base
        else:
            ws.cell(row=i, column=1, value=line).font = title_font if i == 1 else base


def new_book():
    wb = Workbook()
    wb.remove(wb.active)
    return wb
