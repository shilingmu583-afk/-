"""data/*.json から 全国交通事故データ.xlsx を作り直す。

使い方: python3 scripts/build_xlsx.py
"""
import datetime as dt
import json
import pathlib
import urllib.parse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "全国交通事故データ.xlsx"
DOW = "月火水木金土日"
PREFS = ("北海道 青森県 岩手県 宮城県 秋田県 山形県 福島県 茨城県 栃木県 群馬県 埼玉県 千葉県 東京都 神奈川県 "
         "新潟県 富山県 石川県 福井県 山梨県 長野県 岐阜県 静岡県 愛知県 三重県 滋賀県 京都府 大阪府 兵庫県 奈良県 "
         "和歌山県 鳥取県 島根県 岡山県 広島県 山口県 徳島県 香川県 愛媛県 高知県 福岡県 佐賀県 長崎県 熊本県 大分県 "
         "宮崎県 鹿児島県 沖縄県").split()


def load(name):
    return json.loads((ROOT / "data" / name).read_text(encoding="utf-8"))


acc = load("accidents.json")
ref = load("reference.json")
meta = load("meta.json")
for a in acc:
    a["_d"] = dt.date.fromisoformat(a["発生日"])
acc.sort(key=lambda a: (a["_d"], a["都道府県"]))
start = dt.date.fromisoformat(meta["開始日"])
end = dt.date.fromisoformat(meta["収集済み最終日"])

F = "Arial"
hdr_font = Font(name=F, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E78")
base = Font(name=F, size=10)
bold = Font(name=F, bold=True)
link = Font(name=F, size=10, color="0563C1", underline="single")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top")
conf_fill = {"確認済": "E2EFDA", "見出しのみ": "FFF2CC", "要確認": "FCE4D6"}


def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row=1, column=i, value=c)
        x.font, x.fill, x.border = hdr_font, hdr_fill, border
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def hyperlink(url, label):
    return f'=HYPERLINK("{url.replace(chr(34), "%22")}","{label}")'


def gmap(a):
    place = a["場所・道路"]
    q = a["都道府県"] + a["市区町村"]
    if not place.startswith("（"):
        q += " " + place.split("（")[0].replace("・", " ")
    return "https://www.google.com/maps/search/?api=1&query=" + urllib.parse.quote(q)


wb = Workbook()

# 説明
ws0 = wb.active
ws0.title = "説明"
ws0.column_dimensions["A"].width = 110
notes = [
    ("このファイルについて", True),
    (f"対象期間：{start:%Y年%-m月%-d日}〜{end:%Y年%-m月%-d日}の発生分。毎朝自動で前日分を追加しています", False),
    ("収集方法：信頼できる報道機関（NHK・時事通信・地方紙・地方局）と都道府県警の発表を、Web検索で見つかった範囲で収録", False),
    ("", False),
    ("注意", True),
    ("・全件ではありません。全国では1日におよそ5〜10人が交通事故で亡くなっており、ここに載っているのは報道で見つかった一部です", False),
    ("・主に死亡事故です（けがだけの事故はほとんど報道されないため）", False),
    ("・「地図」列は、住所の文字列で Google マップを検索するリンクです。正確な事故地点ではありません", False),
    ("・「緯度」「経度」は町名・交差点が分かるものはその付近、それ以外は市区町村の中心付近の概略値です（「位置の精度」列を参照）。Google マイマップなどに読み込めます", False),
    ("", False),
    ("「確度」列の意味", True),
    ("確認済：報道本文の要約で、場所・日時・内容を確認できたもの", False),
    ("見出しのみ：ニュース一覧の見出しだけで確認したもの。発生日は掲載日から推定。詳細は出典で確認してください", False),
    ("要確認：報道機関以外の出典だけで確認したもの", False),
    ("", False),
    ("シート構成", True),
    ("事故一覧：1行が1件の事故 ／ 都道府県別・日別：事故一覧から自動で集計", False),
    ("参考（道路交通事故以外）：踏切事故・駅構内・駐車場内・作業事故など、警察の交通事故統計に入らないもの", False),
]
for i, (t, b) in enumerate(notes, 1):
    x = ws0.cell(row=i, column=1, value=t)
    x.font = Font(name=F, bold=b, size=12 if b else 10)
    x.alignment = Alignment(wrap_text=True)

# 事故一覧
ws = wb.create_sheet("事故一覧")
cols = ["No", "発生日", "曜日", "都道府県", "市区町村", "場所・道路", "地図", "事故の種類", "死者数", "負傷者数",
        "概要", "確度", "出典", "出典URL", "備考", "緯度", "経度", "位置の精度"]
header(ws, cols, [5, 11, 6, 10, 16, 32, 8, 24, 8, 9, 48, 11, 24, 14, 28, 10, 10, 16])
for n, a in enumerate(acc, 1):
    i = n + 1
    vals = [n, a["_d"], DOW[a["_d"].weekday()], a["都道府県"], a["市区町村"], a["場所・道路"],
            hyperlink(gmap(a), "地図"), a["事故の種類"], a["死者数"], a["負傷者数"], a["概要"], a["確度"],
            a["出典"], hyperlink(a["出典URL"], "記事を開く"), a["備考"], a.get("緯度"), a.get("経度"), a.get("位置の精度", "")]
    for c, v in enumerate(vals, 1):
        x = ws.cell(row=i, column=c, value=v)
        x.font, x.alignment, x.border = base, wrap, border
    ws.cell(row=i, column=2).number_format = "yyyy/mm/dd"
    ws.cell(row=i, column=7).font = ws.cell(row=i, column=14).font = link
    ws.cell(row=i, column=12).fill = PatternFill("solid", fgColor=conf_fill.get(a["確度"], "FFFFFF"))
    for c in (1, 3, 9, 10, 12):
        ws.cell(row=i, column=c).alignment = center
last = len(acc) + 1
tbl = Table(displayName="事故一覧", ref=f"A1:R{last}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
ws.add_table(tbl)
s = last + 2
ws.cell(row=s, column=1, value="合計").font = bold
ws.cell(row=s, column=2, value=f'=COUNTA(D2:D{last})&"件"').font = bold
ws.cell(row=s, column=8, value="死者数計").font = bold
ws.cell(row=s, column=9, value=f"=SUM(I2:I{last})").font = bold
ws.cell(row=s, column=10, value=f"=SUM(J2:J{last})").font = bold
ws.cell(row=s + 1, column=1,
        value="※ 死者数・負傷者数が空欄の行は、報道で人数を確認できなかったもの（合計には含まれない）"
        ).font = Font(name=F, size=9, italic=True)

rng = lambda col: f"事故一覧!${col}$2:${col}${last}"

# 都道府県別
ws2 = wb.create_sheet("都道府県別")
header(ws2, ["都道府県", "件数", "死者数（判明分）"], [12, 8, 16])
for i, p in enumerate(PREFS, 2):
    ws2.cell(row=i, column=1, value=p)
    ws2.cell(row=i, column=2, value=f"=COUNTIF({rng('D')},A{i})")
    ws2.cell(row=i, column=3, value=f"=SUMIF({rng('D')},A{i},{rng('I')})")
e = len(PREFS) + 1
ws2.cell(row=e + 1, column=1, value="合計")
ws2.cell(row=e + 1, column=2, value=f"=SUM(B2:B{e})")
ws2.cell(row=e + 1, column=3, value=f"=SUM(C2:C{e})")

# 日別
ws3 = wb.create_sheet("日別")
header(ws3, ["日付", "曜日", "件数", "死者数（判明分）"], [12, 6, 8, 16])
days = (end - start).days + 1
for k in range(days):
    i, d = k + 2, start + dt.timedelta(days=k)
    ws3.cell(row=i, column=1, value=d).number_format = "yyyy/mm/dd"
    ws3.cell(row=i, column=2, value=DOW[d.weekday()])
    ws3.cell(row=i, column=3, value=f"=COUNTIF({rng('B')},A{i})")
    ws3.cell(row=i, column=4, value=f"=SUMIF({rng('B')},A{i},{rng('I')})")
t = days + 2
ws3.cell(row=t, column=1, value="合計")
ws3.cell(row=t, column=3, value=f"=SUM(C2:C{t - 1})")
ws3.cell(row=t, column=4, value=f"=SUM(D2:D{t - 1})")

for w in (ws2, ws3):
    for row in w.iter_rows(min_row=2):
        for x in row:
            x.font = base
    for x in w[w.max_row]:
        x.font = bold

# 参考
ws4 = wb.create_sheet("参考（道路交通事故以外）")
header(ws4, ["発生日", "都道府県", "市区町村", "場所", "区分", "概要", "出典", "出典URL"],
       [11, 10, 12, 26, 20, 44, 26, 14])
for i, r in enumerate(sorted(ref, key=lambda r: r["発生日"]), 2):
    vals = [dt.date.fromisoformat(r["発生日"]), r["都道府県"], r["市区町村"], r["場所"], r["区分"], r["概要"],
            r["出典"], hyperlink(r["出典URL"], "記事を開く")]
    for c, v in enumerate(vals, 1):
        x = ws4.cell(row=i, column=c, value=v)
        x.font, x.alignment, x.border = base, wrap, border
    ws4.cell(row=i, column=1).number_format = "yyyy/mm/dd"
    ws4.cell(row=i, column=8).font = link

wb.active = 1
wb.save(OUT)
print(f"{OUT.name}: 事故 {len(acc)} 件 / 参考 {len(ref)} 件 / {start}〜{end}")
